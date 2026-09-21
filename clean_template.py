#!/usr/bin/env python3
"""Remove all prior product rows while preserving the Temu workbook structure."""

from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    path = Path(__file__).with_name("Temu_upload.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Template"]
    for row in sheet.iter_rows(min_row=5, max_row=2998, max_col=sheet.max_column):
        for cell in row:
            cell.value = None
            if cell.comment:
                cell.comment = None
    workbook.save(path)


if __name__ == "__main__":
    main()
