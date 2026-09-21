#!/usr/bin/env python3
"""Scrape selected Favorita categories into the supplied Temu upload template."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import logging
import re
import sys
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

from lxml import html as lxml_html
from openpyxl import load_workbook


BASE_URL = "https://www.favorita.bg/"
CATEGORY_URLS = (
    urljoin(BASE_URL, "fc-barcelona.html"),
    urljoin(BASE_URL, "liverpool-fc.html"),
    urljoin(BASE_URL, "real-madrid-cf.html"),
)
USER_AGENT = "Mozilla/5.0 (compatible; FavoritaTemuScraper/1.2)"
TEMPLATE_SHEET = "Template"
FIRST_DATA_ROW = 5
MAX_DATA_ROW = 2998
MAX_IMAGES = 10

# Columns in the supplied Temu template (1-based).
CATEGORY_COL = 5
CATEGORY_NAME_COL = 6
NAME_COL = 12
GOODS_SKU_COL = 13
SKU_COL = 14
ACTION_COL = 15
BRAND_COL = 18
TRADEMARK_COL = 19
DESCRIPTION_COL = 20
DETAIL_IMAGE_COL = 27
VARIATION_THEME_COL = 1008
SIZE_FAMILY_COL = 1009
SUB_SIZE_FAMILY_COL = 1010
SIZE_COL = 1011
COLOR_COL = 1012
MODEL_COL = 1024
SIZE_CHART_METHOD_COL = 1026
SIZE_CHART_UNIT_COL = 1028
SKU_IMAGE_COL = 1093
QTY_COL = 1104
BASE_PRICE_COL = 1105
REFERENCE_LINK_COL = 1106
LIST_PRICE_COL = 1107
WEIGHT_COL = 1109
LENGTH_COL = 1110
WIDTH_COL = 1111
HEIGHT_COL = 1112
SKU_TYPE_COL = 1113
INDIVIDUALLY_PACKED_COL = 1114
PACK_QTY_COL = 1115
PACK_UNIT_COL = 1116
SHIPPING_TEMPLATE_COL = 1124
COUNTRY_COL = 1128
PRODUCT_IDENTIFICATION_COL = 1181
MANUFACTURER_COL = 1182

# The attached template offers this constrained Temu category set. The rules
# below choose the closest valid category available in that file.
CATEGORY_MAP = {
    "БУТИЛКА / ХАЛБА": 34090,
    "ДЕТСКА ТЕНИСКА": 34988,
    "ДЕТСКИ АНЦУГ": 34988,
    "ДЕТСКИ ФУТБОЛЕН ЕКИП": 34988,
    "МЪЖКА ТЕНИСКА": 30469,
    "МЪЖКИ АНЦУГ": 30469,
    "МЪЖКИ СУИЧЪР": 30469,
    "Блуза/Суичър": 30469,
    "НЕСЕСЕР": 31752,
    "ПЛЕТЕНА ШАПКА": 34859,
    "ПОРТМОНЕ / ЧАНТА": 30718,
    "Раница": 31021,
    "РАНИЦА / МЕШКА": 31021,
    "САК": 29158,
    "ФУТБОЛНА ТОПКА": 36731,
    "Хавлия": 11812,
    "ЧАСОВНИК": 31748,
    "ЧАША": 10585,
    "ЧОРАПИ": 34988,
    "ШАЛ": 30253,
    "ШАПКА ИДИОТКА": 34859,
    "ШАПКА С КОЗИРКА": 34859,
    "ЯКЕ": 30469,
    "ВРАТАРСКИ РЪКАВИЦИ": 36854,
    "Зимни ръкавици": 36854,
    "КАЛЕНДАР": 12150,
    "КУТИЯ ЗА ХРАНА": 31752,
    "КЪСИ ПАНТАЛОНИ / БАНСКИ": 30469,
    "КОЛАН": 30625,
    "АКСЕСОАРИ": 30625,
    "БЕБЕШКИ ЕКИП": 34988,
}


@dataclass(frozen=True)
class Variant:
    variant_id: str
    label: str
    price: Decimal
    in_stock: bool
    code: str


@dataclass(frozen=True)
class Product:
    url: str
    team: str
    source_category: str
    name: str
    description: str
    code: str
    images: tuple[str, ...]
    variants: tuple[Variant, ...]
    base_price: Decimal
    in_stock: bool


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    value = html.unescape(value).replace("\xa0", " ")
    return re.sub(r"[ \t]+", " ", re.sub(r"\n[ \t]+", "\n", value)).strip()


def slug_token(value: str) -> str:
    value = clean_text(value).upper()
    value = re.sub(r"[^0-9A-ZА-Я]+", "-", value, flags=re.IGNORECASE)
    return value.strip("-")[:70]


def short_hash(value: str, length: int = 8) -> str:
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:length].upper()


def fit_identifier(value: str, suffix: str = "", max_length: int = 80) -> str:
    value = clean_text(value).strip(" -") or "FAV"
    if not suffix:
        return value[:max_length].rstrip(" -")
    suffix = f"-{suffix.strip(' -')}"
    head = value[: max_length - len(suffix)].rstrip(" -")
    return f"{head}{suffix}"


class Fetcher:
    def __init__(self, retries: int = 3, delay: float = 0.05, timeout: int = 25):
        self.retries = retries
        self.delay = delay
        self.timeout = timeout

    def get(self, url: str) -> str:
        last_error = None
        for attempt in range(1, self.retries + 1):
            try:
                req = Request(url, headers={"User-Agent": USER_AGENT, "Accept-Language": "bg,en;q=0.8"})
                with urlopen(req, timeout=self.timeout) as response:
                    charset = response.headers.get_content_charset() or "utf-8"
                    body = response.read().decode(charset, errors="replace")
                if self.delay:
                    time.sleep(self.delay)
                return body
            except (HTTPError, URLError, TimeoutError, OSError) as exc:
                last_error = exc
                logging.warning("Fetch failed (%s/%s): %s", attempt, self.retries, url)
                if attempt < self.retries:
                    time.sleep(min(1.5 ** attempt, 4))
        raise RuntimeError(f"Unable to download {url}: {last_error}")


def first_text(doc, xpath: str) -> str:
    values = doc.xpath(xpath)
    if not values:
        return ""
    value = values[0]
    if hasattr(value, "text_content"):
        value = value.text_content()
    return clean_text(str(value))


def discover_product_urls(fetcher: Fetcher, category_url: str) -> list[str]:
    found: list[str] = []
    page = 1
    while True:
        url = category_url if page == 1 else f"{category_url}?page={page}"
        doc = lxml_html.fromstring(fetcher.get(url))
        links = [urljoin(BASE_URL, x) for x in doc.xpath('//a[contains(@class,"gs-item-title")]/@href')]
        links = list(dict.fromkeys(links))
        if not links:
            break
        for link in links:
            if link not in found:
                found.append(link)
        next_pages = [int(x) for x in doc.xpath('//a[contains(@class,"gs-paging-link")]/@data-page') if str(x).isdigit()]
        if not next_pages or page >= max(next_pages):
            break
        page += 1
    return found


def parse_variant_json(page_source: str) -> dict:
    match = re.search(r"pub\.product\.variant\.init\('(.*?)',\s*'orderFrm'", page_source, re.S)
    if not match or match.group(1) in ("", "{}"):
        return {}
    encoded = match.group(1)
    try:
        # JavaScript string escaping around JSON is compatible with this targeted unescape.
        decoded = bytes(encoded, "utf-8").decode("unicode_escape")
        return json.loads(decoded)
    except (UnicodeDecodeError, json.JSONDecodeError):
        # Most pages contain only ASCII JSON plus \u escapes; direct parsing covers them.
        return json.loads(encoded)


def parse_price(text: str) -> Decimal:
    match = re.search(r"([0-9]+(?:[.,][0-9]+)?)", clean_text(text))
    if not match:
        raise ValueError(f"Price not found in {text!r}")
    return Decimal(match.group(1).replace(",", ".")).quantize(Decimal("0.01"))


def parse_product(fetcher: Fetcher, url: str) -> Product:
    source = fetcher.get(url)
    doc = lxml_html.fromstring(source)
    name = first_text(doc, '//div[contains(@class,"gs-item-title")]/h1')
    code = first_text(doc, '//strong[@data-code]')
    price_text = first_text(doc, '//strong[@data-var-price]')
    base_price = parse_price(price_text)
    desc_nodes = doc.xpath('//div[contains(@class,"gs-item-txt") and contains(@class,"gs-rtf")]')
    description = clean_text(desc_nodes[0].text_content()) if desc_nodes else first_text(doc, '//meta[@name="description"]/@content')
    # Remove shop-service boilerplate from the end when present.
    description = re.split(r"Допълнителна информация:|Всички поръчки се изпращат|Всички продукти се изпращат", description, maxsplit=1)[0].strip()
    images = []
    for path in doc.xpath('//div[@id="gs-gallery"]//img/@data-zoom-image | //div[@id="gs-gallery"]//img/@src'):
        full = urljoin(BASE_URL, path)
        if full not in images:
            images.append(full)
    team = ""
    source_category = ""
    for row in doc.xpath('//div[contains(@class,"gs-tab-row")]'):
        label = first_text(row, './/strong')
        value = first_text(row, './/span')
        if label.startswith("Отбор"):
            team = value
        elif label.startswith("Категория"):
            source_category = value
    variants = []
    for vid, raw in parse_variant_json(source).items():
        label = clean_text(raw.get("optionsName"))
        price = Decimal(str(raw.get("price", base_price))).quantize(Decimal("0.01"))
        label_says_sold_out = bool(re.search(r"\b(?:изчерпан|изчерпана|sold\s*out|out\s*of\s*stock)\b", label, re.I))
        variants.append(Variant(
            variant_id=str(raw.get("Id", vid)),
            label=label,
            price=price,
            in_stock=bool(raw.get("inStock")) and not label_says_sold_out,
            code=clean_text(raw.get("code")) or code,
        ))
    if variants:
        in_stock = any(v.in_stock for v in variants)
    else:
        availability = first_text(doc, '//meta[@property="og:availability"]/@content').lower()
        button_exists = bool(doc.xpath('//*[@id="addToCartBtn"]'))
        sold_out = bool(doc.xpath('//*[contains(@class,"out-of-stock") or contains(@class,"not-available")]'))
        in_stock = availability in ("instock", "in stock") or (button_exists and not sold_out)
    return Product(
        url=url,
        team=team,
        source_category=source_category,
        name=name,
        description=description,
        code=code,
        images=tuple(images[:MAX_IMAGES]),
        variants=tuple(variants),
        base_price=base_price,
        in_stock=in_stock,
    )


def choose_temu_category(product: Product) -> int:
    title = product.name.lower()
    cat = product.source_category
    if "портф" in title or "портмоне" in title:
        return 30656
    if "подаръчна торб" in title:
        return 17328
    if "чанта за обув" in title:
        return 31111
    if "раница" in title:
        return 31021
    if "сак" in title:
        return 29158
    if "чанта" in title and "обяд" not in title and "термо" not in title:
        return 30718
    if "термо чаша" in title:
        return 12703
    if any(x in title for x in ("чаша", "халба")):
        return 10585
    if "бутилка" in title:
        return 34090
    if "ключодърж" in title:
        return 30625
    if "картичка" in title:
        return 1111
    if "бадж" in title or "значк" in title:
        return 1432
    if "молив" in title:
        return 1291
    if "химикал" in title:
        return 1180
    if "шал" in title:
        return 30253
    if "топка" in title:
        return 36731
    if "вратарски ръкавиц" in title:
        return 36854
    if "хавли" in title:
        return 11812
    if "чадър" in title:
        return 31015
    if "пижам" in title:
        return 30827
    if "часовник" in title:
        return 31748
    if "шапка" in title:
        return 34859
    if any(x in title for x in ("детск", "yamal", "salah", "bellingham", "mbapp")) and any(
        x in title for x in ("екип", "тениска", "анцуг", "худи", "ветровка", "яке", "бански", "чорап")
    ):
        return 34988
    return CATEGORY_MAP.get(cat, 31752)


def extract_color(name: str, description: str) -> str:
    text = f"{name} {description}".lower()
    colors = (
        ("Тъмносиня", ("тъмносиня", "тъмносин", "navy")),
        ("Черна", ("черна", "черен", "black")), ("Бяла", ("бяла", "бял", "white")),
        ("Червена", ("червена", "червен", "red")), ("Синя", ("синя", "син", "blue")),
        ("Зелена", ("зелена", "зелен")),
        ("Сива", ("сива", "сив")), ("Жълта", ("жълта", "жълт")),
    )
    for label, needles in colors:
        if any(x in text for x in needles):
            return label
    return ""


def split_variant_label(label: str) -> tuple[str, str]:
    """Return a clean size and, when encoded after ';', a color/finish label."""
    label = re.sub(r"^ИЗЧЕРПАН\s*", "", clean_text(label), flags=re.I)
    if ";" not in label:
        return label, ""
    size, detail = (clean_text(part) for part in label.split(";", 1))
    detail = re.sub(r"\s*\([^)]*%[^)]*\)\s*$", "", detail).strip(" -–")
    detail = re.sub(r"^Комплект\s*[-–]\s*", "", detail, flags=re.I)
    return size, detail[:80]


def size_profile(category_id: int, size_label: str) -> tuple[str, str]:
    label = clean_text(size_label)
    if not label or re.search(r"one\s*size|един\s*размер", label, re.I):
        return "2 - Regular Size", "1 - One Size"
    if category_id == 36731:
        return "101 - Custom size", "10 - Alpha"
    if re.search(r"(?:г\.|год|ръст|cm|см)", label, re.I):
        if category_id in (34988, 30827):
            return "2 - Regular Size", "8 - Age"
        return "101 - Custom size", "10 - Alpha"
    if re.fullmatch(r"(?:EU\s*)?\d+(?:[-/]\d+)?", label, re.I):
        return "2 - Regular Size", "7 - Numeric"
    if re.fullmatch(r"(?:X{0,4}[SLM]|[2-6]XL)(?:[-/](?:X{0,4}[SLM]|[2-6]XL))?", label, re.I):
        return "2 - Regular Size", "10 - Alpha"
    return "101 - Custom size", "10 - Alpha"


def team_brand(team: str, name: str) -> tuple[str, str]:
    text = f"{team} {name}".lower()
    if "barcelona" in text:
        return "FC BARCELONA", "FC Barcelona"
    if "liverpool" in text:
        return "LIVERPOOL FC", "Liverpool FC"
    return "REAL MADRID", "Real Madrid"


def team_country(team: str, name: str) -> str:
    return "United Kingdom" if "liverpool" in f"{team} {name}".lower() else "Spain"


def infer_dimensions(description: str, category_id: int) -> tuple[int, int, int, int]:
    weight_match = re.search(r"(?:Тегло|Weight)\s*:?\s*(\d+(?:[.,]\d+)?)\s*(kg|кг|g|гр)", description, re.I)
    parsed_weight = None
    if weight_match:
        parsed_weight = Decimal(weight_match.group(1).replace(",", "."))
        if weight_match.group(2).lower() in ("kg", "кг"):
            parsed_weight *= 1000
        parsed_weight = max(1, int(parsed_weight.quantize(Decimal("1"))))
    match = re.search(r"(?:Размер(?:и)?|Dimensions?)\s*:?\s*(\d+(?:[.,]\d+)?)\s*[xх×]\s*(\d+(?:[.,]\d+)?)(?:\s*[xх×]\s*(\d+(?:[.,]\d+)?))?\s*см", description, re.I)
    if match:
        vals = [max(1, int(Decimal(x.replace(",", ".")).quantize(Decimal("1")))) for x in match.groups() if x]
        while len(vals) < 3:
            vals.append(5)
        return parsed_weight or category_defaults(category_id)[0], vals[0], vals[1], vals[2]
    weight, length, width, height = category_defaults(category_id)
    return parsed_weight or weight, length, width, height


def category_defaults(category_id: int) -> tuple[int, int, int, int]:
    defaults = {
        1111: (50, 20, 15, 1), 1180: (30, 15, 3, 2), 1291: (150, 22, 15, 3),
        1432: (50, 10, 8, 2), 10585: (400, 12, 10, 10), 11812: (500, 35, 25, 8),
        12150: (200, 25, 15, 10), 12703: (450, 22, 10, 10), 17328: (100, 35, 25, 5),
        29158: (600, 40, 30, 15), 30253: (250, 30, 25, 5), 30469: (350, 32, 25, 5),
        30625: (80, 12, 8, 3), 30656: (250, 15, 12, 5), 30718: (500, 35, 25, 12),
        30827: (450, 35, 28, 7), 31015: (450, 90, 8, 8), 31021: (700, 45, 32, 18),
        31111: (200, 42, 32, 5), 31748: (250, 15, 12, 8), 31752: (500, 35, 28, 12),
        34090: (350, 28, 10, 10), 34859: (180, 25, 20, 10), 34988: (350, 32, 25, 6),
        36731: (450, 23, 23, 23), 36854: (350, 30, 18, 10),
    }
    return defaults.get(category_id, (500, 30, 25, 10))


def header_columns(ws) -> dict[str, list[int]]:
    columns: dict[str, list[int]] = defaultdict(list)
    for col in range(1, ws.max_column + 1):
        key = ws.cell(4, col).value
        if key:
            columns[str(key)].append(col)
    return dict(columns)


def dropdown_index(wb) -> dict[str, tuple[str, ...]]:
    ws = wb["Dropdown Lists"]
    result = {}
    for row in range(1, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if not key:
            continue
        result[str(key)] = tuple(
            str(ws.cell(row, col).value)
            for col in range(3, ws.max_column + 1)
            if ws.cell(row, col).value not in (None, "")
        )
    return result


def set_first(ws, row: int, columns: dict[str, list[int]], key: str, value) -> None:
    if value not in (None, "") and columns.get(key):
        ws.cell(row, columns[key][0], value)


def pick_allowed(choices: tuple[str, ...], candidates: Iterable[str]) -> str:
    if not choices:
        return ""
    normalized = {clean_text(choice).casefold(): choice for choice in choices}
    for candidate in candidates:
        match = normalized.get(clean_text(candidate).casefold())
        if match:
            return match
    for candidate in candidates:
        needle = clean_text(candidate).casefold()
        for choice in choices:
            if needle and needle in clean_text(choice).casefold():
                return choice
    return choices[0]


def property_choices(dropdowns: dict[str, tuple[str, ...]], category_id: int, label: str) -> tuple[str, ...]:
    return dropdowns.get(f"t_3_{category_id}_{label}", ())


def material_candidates(product: Product, category_id: int) -> list[str]:
    text = f"{product.name} {product.description}".lower()
    rules = (
        (("естествена кожа", "телешка кожа"), ["Top Layer Cowhide", "Genuine Leather", "Top grain leather", "Leather"]),
        (("изкуствена кожа", "еко кожа"), ["Faux leather", "PU Leather", "Synthetic Leather", "PU"]),
        (("порцелан",), ["Porcelain", "Ceramics", "Ceramic"]),
        (("керами",), ["Ceramics", "Ceramic"]),
        (("неръждаем",), ["304 Stainless Steel", "Stainless Steel"]),
        (("алумини",), ["Aluminum Alloy", "Aluminum"]),
        (("100% памук", "памук"), ["Cotton", "Cotton Blend"]),
        (("полиестер",), ["Polyester", "Polyester Blend", "Polyester (polyester Fiber)"]),
        (("акрил",), ["Acrylic"]), (("найлон",), ["Nylon"]), (("силикон",), ["Silicone"]),
        (("пластмас",), ["Plastic", "PP (polypropylene)", "Polypropylene"]),
        (("харт",), ["Paper"]), (("метал",), ["Metal", "Iron", "Stainless Steel"]),
        (("стъкло",), ["Glass", "Ordinary Glass"]), (("дърво",), ["Wood"]),
        (("гума", "каучук"), ["Rubber"]), ((" pu ", "полиуретан"), ["PU", "Polyurethane"]),
    )
    for needles, candidates in rules:
        if any(needle in f" {text} " for needle in needles):
            return candidates
    fallbacks = {
        1111: ["Paper"], 1180: ["Plastic"], 1291: ["Wood"], 1432: ["Metal"],
        10585: ["Ceramic"], 11812: ["Cotton"], 12150: ["Paper", "Resin"],
        12703: ["Stainless Steel"], 17328: ["Paper"], 29158: ["Polyester", "Fabric"],
        30253: ["Polyester"], 30469: ["Cotton", "Polyester"], 30625: ["Metal"],
        30656: ["Faux leather"], 30718: ["Polyester", "Fabric"], 30827: ["Cotton"],
        31021: ["Polyester"], 31111: ["Polyester"], 31748: ["Metal"],
        34090: ["Aluminum Alloy"], 34859: ["Polyester"], 34988: ["Polyester"],
        36731: ["PU", "PVC"], 36854: ["Polyester"],
    }
    return fallbacks.get(category_id, ["Polyester", "Plastic", "Metal"])


def set_choice_property(ws, row: int, columns: dict[str, list[int]], dropdowns, category_id: int, key: str, label: str, candidates: Iterable[str]) -> None:
    value = pick_allowed(property_choices(dropdowns, category_id, label), candidates)
    set_first(ws, row, columns, key, value)


def inferred_composition(product: Product, category_id: int) -> tuple[str, int]:
    text = f"{product.name} {product.description}".lower()
    materials = (
        ("Cotton", "памук"), ("Polyester", "полиестер"), ("Acrylic", "акрил"),
        ("Wool", "вълна"), ("Polyamide", "полиамид"), ("Elastane", "еластан"),
    )
    for material, needle in materials:
        if needle in text:
            nearby = re.search(rf"(\d{{1,3}})\s*%\s*{needle}|{needle}[^0-9]{{0,15}}(\d{{1,3}})\s*%", text)
            if nearby:
                pct = int(next(x for x in nearby.groups() if x is not None))
                return material, min(100, max(1, pct))
            return material, 100
    return ("Cotton", 100) if category_id in (30469, 30827, 11812) else ("Polyester", 100)


def set_composition(ws, row: int, columns: dict[str, list[int]], group: int, material: str, percent: int) -> None:
    codes = {"Acrylic": 74, "Cotton": 78, "Polyamide": 97, "Polyester": 98, "Wool": 110, "Elastane": 35353}
    code = codes.get(material, 98)
    set_first(ws, row, columns, f"t_3_Property:{group}:{code}", percent)


def fill_category_attributes(wb, ws, row: int, columns: dict[str, list[int]], dropdowns, category_id: int, product: Product) -> None:
    candidates = material_candidates(product, category_id)
    material12 = {1432, 30469, 30625, 30656, 30718, 30827, 31021, 31111, 34859, 34988, 36731}
    material121 = {10585, 1180, 11812, 12150, 12703, 17328, 29158}
    if category_id in material12:
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:12", "12 - Material", candidates)
    if category_id in material121:
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:121", "121 - Material", candidates)

    composition_categories = {30253, 30469, 34859, 34988, 11812}
    if category_id in composition_categories:
        material, percent = inferred_composition(product, category_id)
        set_composition(ws, row, columns, 15, material, percent)
    if category_id == 30827:
        material, percent = inferred_composition(product, category_id)
        set_composition(ws, row, columns, 1428, material, percent)
        set_composition(ws, row, columns, 1429, material, percent)

    age_candidates = {
        1180: ["6 Years+"], 1291: ["3 Years+"], 30827: ["12 and under"],
        31021: ["3 Years And Older (not Including 3 Years Old)"], 34988: ["12 and under"],
        36731: ["6 Years+"], 36854: ["6 Years+"],
    }
    if category_id in age_candidates:
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:1117", "1117 - Applicable Age Group", age_candidates[category_id])

    for cat in (30718, 31021, 34090):
        if category_id == cat:
            set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:1067", "1067 - Power Mode", ["Without electricity"])

    if category_id in (31748, 34090, 1111):
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:1920", "1920 - Major Material", candidates)
    if category_id in (10585, 12703, 34090):
        food_candidates = material_candidates(product, category_id) + ["Ceramics", "Aluminum", "Polypropylene (PP)"]
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:8319", "8319 - Food Contact Material", food_candidates)
    if category_id == 34090:
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:3989", "3989 - Liner Material", candidates)
    if category_id == 12703:
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:1561", "1561 - Power Supply", ["Use Without Electricity"])
    if category_id == 17328:
        value = pick_allowed(property_choices(dropdowns, category_id, "3980 - Square Gram Weight (g/㎡)"), ["≥86＜200", "<86"])
        if value and columns.get("t_3_Property:3980"):
            ws.cell(row, columns["t_3_Property:3980"][-1], value)
    if category_id == 11812:
        set_first(ws, row, columns, "t_3_Property:3980", 400)
    if category_id == 36854:
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:2421", "2421 - Fabric Material", candidates)
        has_guards = "Yes" if re.search(r"finger\s*(?:save|guard)|протектор|шини", product.description, re.I) else "No"
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:6921", "6921 - Does It Have Finger Guards?", [has_guards])
    if category_id == 31015:
        opening = "Automatic" if re.search(r"автомат", product.description, re.I) else "Manual"
        ribs = re.search(r"(\d+)\s*(?:ребра|спици|ribs)", product.description, re.I)
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:1224", "1224 - Open Way", [opening])
        set_choice_property(ws, row, columns, dropdowns, category_id, "t_3_Property:2321", "2321 - Number Of Ribs", [f"{ribs.group(1) if ribs else 8} Ribs"])


def estimated_garment_measurements(size_label: str) -> tuple[int, int, int, int]:
    label = clean_text(size_label).upper()
    alpha = {
        "XS": (90, 67, 74, 100), "S": (96, 69, 78, 102), "M": (102, 71, 84, 104),
        "L": (108, 73, 90, 106), "XL": (114, 75, 96, 108), "XXL": (120, 77, 102, 110),
        "3XL": (126, 79, 108, 112), "4XL": (132, 81, 114, 114),
    }
    for key in sorted(alpha, key=len, reverse=True):
        if re.search(rf"\b{re.escape(key)}\b", label):
            return alpha[key]
    height_match = re.search(r"(9[0-9]|1[0-8][0-9])\s*(?:CM|СМ)", label)
    height = int(height_match.group(1)) if height_match else 140
    chest = max(56, round(height * 0.54))
    length = max(40, round(height * 0.43))
    waist = max(52, round(height * 0.48))
    pants = max(55, round(height * 0.63))
    return chest, length, waist, pants


def fill_size_chart(ws, row: int, columns: dict[str, list[int]], category_id: int, size_label: str) -> None:
    if category_id not in (11812, 30469, 30827, 34988):
        return
    set_first(ws, row, columns, "t_5_Size Chart Method", "Add size chart manually")
    set_first(ws, row, columns, "t_5_Unit", "cm-g-ml")
    chest, length, waist, pants = estimated_garment_measurements(size_label)
    if category_id == 30469:
        set_first(ws, row, columns, "t_5_Size Chart Element:3:Product:10002", chest)
        set_first(ws, row, columns, "t_5_Size Chart Element:3:Product:10003", length)
    elif category_id == 30827:
        for key, value in (
            ("t_5_Size Chart Element:128:Product:10002", chest),
            ("t_5_Size Chart Element:128:Product:10003", length),
            ("t_5_Size Chart Element:129:Product:10002", chest),
            ("t_5_Size Chart Element:129:Product:10003", length + 25),
            ("t_5_Size Chart Element:130:Product:10002", chest),
            ("t_5_Size Chart Element:130:Product:10003", length),
            ("t_5_Size Chart Element:127:Product:10005", waist),
            ("t_5_Size Chart Element:127:Product:10008", pants),
        ):
            set_first(ws, row, columns, key, value)


def effective_variant_price(product: Product, variant: Variant) -> Decimal:
    if variant.price > 0:
        return variant.price
    _, detail = split_variant_label(variant.label)
    candidates = []
    for sibling in product.variants:
        if not sibling.in_stock or sibling.price <= 0:
            continue
        _, sibling_detail = split_variant_label(sibling.label)
        if detail and sibling_detail.casefold() == detail.casefold():
            candidates.append(sibling.price)
    if candidates:
        return Counter(candidates).most_common(1)[0][0]
    if product.base_price > 0:
        return product.base_price
    candidates = [v.price for v in product.variants if v.in_stock and v.price > 0]
    if candidates:
        return Counter(candidates).most_common(1)[0][0]
    raise ValueError(f"No positive price for {product.url} / {variant.label}")


def build_parent_skus(products: Iterable[Product]) -> dict[str, str]:
    groups: dict[str, list[Product]] = defaultdict(list)
    seen_urls = set()
    for product in products:
        if product.url in seen_urls:
            continue
        seen_urls.add(product.url)
        base = product.code or f"FAV-{slug_token(product.name)}"
        groups[base].append(product)
    result = {}
    used = set()
    for base, group in groups.items():
        unique_urls = {p.url for p in group}
        for product in group:
            suffix = short_hash(product.url) if len(unique_urls) > 1 else ""
            candidate = fit_identifier(base, suffix)
            if candidate in used:
                candidate = fit_identifier(base, short_hash(product.url + product.name, 10))
            used.add(candidate)
            result[product.url] = candidate
    return result


def clear_template_rows(ws) -> None:
    # Only touch cells that already exist in the workbook. Iterating the full
    # 2,994 x 1,182 product area creates more than 3.5 million Python objects
    # and is the main reason GitHub Actions runs used to take over an hour.
    for (row, _), cell in list(ws._cells.items()):
        if row >= FIRST_DATA_ROW:
            cell.value = None
            if cell.comment:
                cell.comment = None


def formula_for_category_name(row: int) -> str:
    return '=IFERROR(VLOOKUP(INDIRECT(ADDRESS(ROW(), COLUMN()-1)), \'Category Name\'!$A:$B, 2, FALSE), "Please select the [category] first. After selecting, the full category path will be displayed here. Please do not modify the functions in the cells")'


def product_rows(products: Iterable[Product]):
    for product in products:
        if product.variants:
            for variant in product.variants:
                if variant.in_stock:
                    yield product, variant
        elif product.in_stock:
            yield product, None


def validate_output(wb, row_count: int) -> None:
    ws = wb[TEMPLATE_SHEET]
    end_row = FIRST_DATA_ROW + row_count - 1
    rows = range(FIRST_DATA_ROW, end_row + 1)
    columns = header_columns(ws)

    def first(key: str) -> int:
        return columns[key][0]

    contribution_goods = first("t_1_Contribution Goods")
    contribution_sku = first("t_1_Contribution SKU")
    reference_link = first("t_6_Reference Link")
    base_price = first("t_6_Base Price - EUR")
    list_price = first("t_6_List Price - EUR")
    image = first("t_6_SKU Images URL")

    sku_values = [ws.cell(row, contribution_sku).value for row in rows]
    duplicates = [value for value, count in Counter(sku_values).items() if value and count > 1]
    if duplicates:
        raise RuntimeError(f"Duplicate Contribution SKU values: {duplicates[:10]}")

    parent_urls: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        parent_urls[str(ws.cell(row, contribution_goods).value)].add(str(ws.cell(row, reference_link).value))
    collisions = {parent: urls for parent, urls in parent_urls.items() if len(urls) > 1}
    if collisions:
        raise RuntimeError(f"Contribution Goods values merge different products: {list(collisions)[:10]}")

    for row in rows:
        price = ws.cell(row, base_price).value
        comparison = ws.cell(row, list_price).value
        if not isinstance(price, (int, float)) or price <= 0:
            raise RuntimeError(f"Non-positive Base Price at row {row}: {price}")
        if not isinstance(comparison, (int, float)) or abs(comparison - price * 2) > 0.001:
            raise RuntimeError(f"Incorrect List Price at row {row}: {comparison}")
        if not ws.cell(row, image).value:
            raise RuntimeError(f"Missing SKU image at row {row}")

    dropdowns = dropdown_index(wb)
    category_col = first("t_1_Category")
    theme_col = first("t_4_Variation Theme")
    goods_mode = wb["GoodsLevelMode"]
    mode_rows = {
        str(goods_mode.cell(row, 1).value): row
        for row in range(1, goods_mode.max_row + 1)
        if goods_mode.cell(row, 1).value
    }

    def requirement_group(key: str) -> str:
        if key.startswith("t_3_Property:"):
            parts = key.split(":")
            if len(parts) >= 3:
                return ":".join(parts[:2])
        return key

    required_groups: dict[int, dict[str, list[int]]] = {}

    for row in rows:
        category_id = int(ws.cell(row, category_col).value)
        theme = ws.cell(row, theme_col).value
        allowed = dropdowns.get(f"t_4_{category_id}_Variation Theme", ())
        if allowed and theme not in allowed:
            raise RuntimeError(f"Invalid Variation Theme at row {row}: {theme!r} for {category_id}")

        groups = required_groups.get(category_id)
        if groups is None:
            require_row = mode_rows.get(f"{category_id}_require")
            if not require_row:
                raise RuntimeError(f"Missing template requirement definition for category {category_id}")
            groups = defaultdict(list)
            for col in range(1, goods_mode.max_column + 1):
                if str(goods_mode.cell(require_row, col).value).lower() != "require":
                    continue
                key = str(ws.cell(4, col).value)
                if key == "t_8_Governance Property:2":
                    continue  # Intentionally left blank by merchant instruction.
                groups[requirement_group(key)].append(col)
            required_groups[category_id] = groups
        for key, required_columns in groups.items():
            if all(ws.cell(row, col).value in (None, "") for col in required_columns):
                label = ws.cell(2, required_columns[0]).value
                raise RuntimeError(f"Missing required field at row {row}: {label} ({key})")


def write_output(template: Path, output: Path, products: list[Product]) -> int:
    wb = load_workbook(template)
    ws = wb[TEMPLATE_SHEET]
    clear_template_rows(ws)
    rows = list(product_rows(products))
    if len(rows) > MAX_DATA_ROW - FIRST_DATA_ROW + 1:
        raise RuntimeError(f"Too many output rows ({len(rows)}) for the template")

    columns = header_columns(ws)
    dropdowns = dropdown_index(wb)
    parent_skus = build_parent_skus(product for product, _ in rows)
    used_child_skus: set[str] = set()

    for offset, (product, variant) in enumerate(rows):
        row = FIRST_DATA_ROW + offset
        category_id = choose_temu_category(product)
        price = effective_variant_price(product, variant) if variant else product.base_price
        if price <= 0:
            raise ValueError(f"No positive price for {product.url}")
        price = price.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        list_price = (price * 2).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        brand, trademark = team_brand(product.team, product.name)
        variant_label = variant.label if variant else ""
        parent_sku = parent_skus[product.url]
        if variant:
            suffix = f"{slug_token(variant_label)[:42]}-{variant.variant_id}"
            child_sku = fit_identifier(parent_sku, suffix)
        else:
            child_sku = parent_sku
        if child_sku in used_child_skus:
            child_sku = fit_identifier(parent_sku, short_hash(f"{product.url}|{variant.variant_id if variant else ''}", 10))
        used_child_skus.add(child_sku)
        weight, length, width, height = infer_dimensions(product.description, category_id)

        raw_size, variant_color = split_variant_label(variant_label)
        product_color = extract_color(product.name, product.description)
        if category_id in (30469, 30827, 34988):
            variation_theme = "Color × Size"
            size_value = raw_size or "One size"
            color_value = variant_color or extract_color(variant_label, "") or product_color or "Многоцветна"
        elif category_id == 36731:
            variation_theme = "Size"
            ball_size = re.search(r"(?:размер|size)\s*:?\s*(\d+)", product.description, re.I)
            size_value = raw_size or (f"Size {ball_size.group(1)}" if ball_size else "One size")
            color_value = product_color or None
        elif variant_label:
            variation_theme = "Size"
            size_value = variant_label
            color_value = variant_color or product_color or None
        else:
            variation_theme = "Model"
            size_value = ""
            color_value = product_color or None
        size_family, sub_size_family = size_profile(category_id, size_value)

        ws.cell(row, CATEGORY_COL, category_id)
        ws.cell(row, CATEGORY_NAME_COL, formula_for_category_name(row))
        ws.cell(row, NAME_COL, product.name[:500])
        ws.cell(row, GOODS_SKU_COL, parent_sku)
        ws.cell(row, SKU_COL, child_sku)
        ws.cell(row, ACTION_COL, "Add")
        ws.cell(row, BRAND_COL, brand)
        ws.cell(row, TRADEMARK_COL, trademark)
        ws.cell(row, DESCRIPTION_COL, product.description[:5000])
        for idx, image in enumerate(product.images[:MAX_IMAGES]):
            ws.cell(row, DETAIL_IMAGE_COL + idx, image)
            ws.cell(row, SKU_IMAGE_COL + idx, image)
        ws.cell(row, VARIATION_THEME_COL, variation_theme)
        ws.cell(row, SIZE_FAMILY_COL, size_family if size_value else None)
        ws.cell(row, SUB_SIZE_FAMILY_COL, sub_size_family if size_value else None)
        ws.cell(row, SIZE_COL, size_value or None)
        ws.cell(row, COLOR_COL, color_value)
        ws.cell(row, MODEL_COL, product.team or brand)
        fill_category_attributes(wb, ws, row, columns, dropdowns, category_id, product)
        fill_size_chart(ws, row, columns, category_id, size_value)
        ws.cell(row, QTY_COL, 10)
        ws.cell(row, BASE_PRICE_COL, float(price))
        ws.cell(row, REFERENCE_LINK_COL, product.url)
        ws.cell(row, LIST_PRICE_COL, float(list_price))
        ws.cell(row, WEIGHT_COL, weight)
        ws.cell(row, LENGTH_COL, length)
        ws.cell(row, WIDTH_COL, width)
        ws.cell(row, HEIGHT_COL, height)
        ws.cell(row, SKU_TYPE_COL, "Single set")
        ws.cell(row, INDIVIDUALLY_PACKED_COL, "Yes")
        ws.cell(row, PACK_QTY_COL, 1)
        ws.cell(row, PACK_UNIT_COL, "piece")
        ws.cell(row, SHIPPING_TEMPLATE_COL, "Доставка")
        ws.cell(row, COUNTRY_COL, team_country(product.team, product.name))
        ws.cell(row, PRODUCT_IDENTIFICATION_COL, child_sku)
        ws.cell(row, MANUFACTURER_COL, trademark)

    validate_output(wb, len(rows))
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return len(rows)


def run(
    template: Path,
    output: Path,
    limit: int = 0,
    workers: int = 8,
    retries: int = 3,
    timeout: int = 25,
    delay: float = 0.05,
) -> None:
    fetcher = Fetcher(retries=retries, delay=delay, timeout=timeout)
    all_urls: list[str] = []
    for category_url in CATEGORY_URLS:
        links = discover_product_urls(fetcher, category_url)
        logging.info("Found %d products in %s", len(links), category_url)
        all_urls.extend(links)
    all_urls = list(dict.fromkeys(all_urls))
    if limit:
        all_urls = all_urls[:limit]
    product_by_url = {}
    failures = []
    completed = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        futures = {executor.submit(parse_product, fetcher, url): url for url in all_urls}
        for future in as_completed(futures):
            url = futures[future]
            completed += 1
            try:
                product = future.result()
                product_by_url[url] = product
                logging.info("[%d/%d] %s", completed, len(all_urls), product.name)
            except Exception as exc:
                failures.append((url, str(exc)))
                logging.error("Failed product %s: %s", url, exc)
    if failures:
        sample = ", ".join(url for url, _ in failures[:5])
        raise RuntimeError(f"Failed to download {len(failures)} products after retries: {sample}")
    products = [product_by_url[url] for url in all_urls if url in product_by_url]
    if not products:
        raise RuntimeError("No products were successfully scraped")
    rows = write_output(template, output, products)
    logging.info("Saved %d in-stock SKU rows to %s", rows, output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=Path(__file__).with_name("Temu_upload.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("output") / "favorita_temu_upload.xlsx")
    parser.add_argument("--limit", type=int, default=0, help="Testing only: scrape at most N products")
    parser.add_argument("--workers", type=int, default=8, help="Number of concurrent product downloads")
    parser.add_argument("--retries", type=int, default=3, help="Download attempts per page")
    parser.add_argument("--timeout", type=int, default=25, help="Read timeout per download attempt in seconds")
    parser.add_argument("--delay", type=float, default=0.05, help="Polite delay after each successful request")
    parser.add_argument("--verbose", action="store_true")
    args = parser.parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    try:
        run(args.template, args.output, args.limit, args.workers, args.retries, args.timeout, args.delay)
        return 0
    except Exception as exc:
        logging.exception("Scraper failed: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
